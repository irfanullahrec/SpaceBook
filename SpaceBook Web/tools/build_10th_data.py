# Script to generate datasets for Class 10th missing subjects
import docx
import json
import os
import sys

def build_cs_10():
    folder = r"D:\SpaceBook\Books\10th\10th CS\Word"
    paras = []
    for f in sorted(os.listdir(folder)):
        if f.endswith(".docx"):
            doc = docx.Document(os.path.join(folder, f))
            paras.extend([p.text.strip() for p in doc.paragraphs if p.text.strip()])
            
    bounds = [
        (1, "Programming Techniques", "پروگرامنگ کا طریقہ کار", 71, 950),
        (2, "Programming in C", "C میں پروگرامنگ", 950, 2000),
        (3, "Input and Output Handling", "ان پٹ اور آؤٹ پٹ کا عمل", 2000, 3050),
        (4, "Control Structures", "کنٹرول سٹرکچر", 3050, 3700),
        (5, "Loop Structures", "لوپ سٹرکچر", 3700, 4299),
        (6, "Computer Logic and Logic Gates", "کمپیوٹر لاجک اور لاجک گیٹس", 4299, 5270),
        (7, "World Wide Web and HTML", "ورلڈ وائیڈ ویب اور HTML", 5270, len(paras))
    ]
    
    units = []
    for num, title, titleUr, start, end in bounds:
        u_paras = paras[start:end]
        clean_p = [p for p in u_paras if not any(w in p.lower() for w in ["awaz e inqilab", "camscanner", "download pdf", "whatsapp", "not for sale", "maktaba tul ishaat"]) and len(p) > 2]
        sec_paras = clean_p[:40]
        full_text = "\n\n".join(sec_paras[:15])
        
        unit_obj = {
            "number": num,
            "id": f"unit-{num}",
            "title": title,
            "titleUrdu": titleUr,
            "type": "chapter",
            "sections": [
                {
                    "heading": "Core Concepts & Programming Practice",
                    "headingUrdu": "بنیادی تصورات و عملی کوڈنگ",
                    "text": full_text,
                    "paras": sec_paras[:12]
                }
            ],
            "exercise": {
                "shortQuestions": [
                    {"q": f"Explain the main purpose of {title}.", "a": f"{title} teaches essential computer science foundations, algorithm design, and coding constructs in accordance with the 10th grade syllabus."},
                    {"q": f"List key programming concepts covered in Unit {num}.", "a": "1. Syntax, semantics, and standard operations.\n2. Logic structure and computational efficiency.\n3. Practical applications in problem-solving."}
                ],
                "longQuestions": [
                    {"q": f"Write a comprehensive explanation of {title} with suitable examples.", "a": f"Unit {num} covers comprehensive principles of {title}, demonstrating how algorithms, flowcharts, data handling, and program logic function together."}
                ],
                "mcqs": [
                    {"q": f"What is the primary topic of Unit {num} ({title})?", "options": ["Computer programming and logic constructs", "Artistic drawing", "Human anatomy", "Geographical regions"], "answer": 0, "exp": f"Unit {num} focuses on {title} for Class 10."},
                    {"q": "A sequence of step-by-step instructions to solve a problem is called:", "options": ["Algorithm", "Hardware", "Compiler", "Operating System"], "answer": 0, "exp": "An algorithm is a finite sequence of well-defined computer-implementable instructions."}
                ]
            },
            "englishSummary": f"Unit {num} introduces {title}, focusing on programming logic, structured coding, and computational problem solving.",
            "urduSummary": f"یہ یونٹ {titleUr} کے بنیادی تصورات، پروگرامنگ قواعد اور عملی کوڈنگ کی تفصیلی رہنمائی فراہم کرتا ہے۔"
        }
        units.append(unit_obj)
        
    out_file = r"D:\SpaceBook\SpaceBook Web\js\computer_10_data.js"
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Class 10 Computer Science Dataset (KPTBB)\n")
        f.write("const COMPUTER_10_DATA = " + json.dumps(units, ensure_ascii=False, indent=2) + ";\n\n")
        f.write("if (typeof DATA !== 'undefined' && DATA) { DATA.comp10Chapters = COMPUTER_10_DATA; }\n")
        f.write("if (typeof window !== 'undefined') { window.COMPUTER_10_DATA = COMPUTER_10_DATA; }\n")
    print("Generated computer_10_data.js!")

def build_gen_science_10():
    folder = r"D:\SpaceBook\Books\10th\10th General Science\Word"
    paras = []
    for f in sorted(os.listdir(folder)):
        if f.endswith(".docx"):
            doc = docx.Document(os.path.join(folder, f))
            paras.extend([p.text.strip() for p in doc.paragraphs if p.text.strip()])
            
    # Units 6 through 11:
    bounds = [
        (6, "Electricity in Everyday Life", "روزمرہ زندگی میں بجلی", 72, 800),
        (7, "Chemical Reactions & Practical Applications", "کیمیائی تعاملات اور عملی اطلاقات", 800, 1600),
        (8, "Biotechnology", "بائیو ٹیکنالوجی", 1600, 2400),
        (9, "Water Resources", "پانی کے وسائل", 2400, 3200),
        (10, "Environmental Problems & Management", "ماحولیاتی مسائل اور ان کا حل", 3200, 3900),
        (11, "Science, Technology & Development", "سائنس، ٹیکنالوجی اور ترقی", 3900, len(paras))
    ]
    
    units = []
    for num, title, titleUr, start, end in bounds:
        u_paras = paras[start:end]
        clean_p = [p for p in u_paras if not any(w in p.lower() for w in ["awaz e inqilab", "camscanner", "download pdf", "whatsapp", "not for sale"]) and len(p) > 2]
        sec_paras = clean_p[:40]
        full_text = "\n\n".join(sec_paras[:15])
        
        unit_obj = {
            "number": num,
            "id": f"unit-{num}",
            "title": title,
            "titleUrdu": titleUr,
            "type": "chapter",
            "sections": [
                {
                    "heading": "Key Scientific Concepts & Applications",
                    "headingUrdu": "اہم سائنسی تصورات و اطلاقات",
                    "text": full_text,
                    "paras": sec_paras[:12]
                }
            ],
            "exercise": {
                "shortQuestions": [
                    {"q": f"State the primary objective of Unit {num} ({title}).", "a": f"This unit explores scientific principles, empirical mechanisms, and societal applications related to {title}."},
                    {"q": f"List two major takeaways from {title}.", "a": "1. Core scientific explanations and definitions.\n2. Real-world importance and modern environmental or technological impacts."}
                ],
                "longQuestions": [
                    {"q": f"Provide an in-depth analysis of {title} with practical examples.", "a": f"Unit {num} provides comprehensive theoretical and practical knowledge on {title}, fostering scientific awareness and environmental responsibility."}
                ],
                "mcqs": [
                    {"q": f"What is the central focus of Unit {num} ({title})?", "options": ["Scientific principles and applications", "Classical literature", "Grammar rules", "Ancient poetry"], "answer": 0, "exp": f"Unit {num} covers scientific principles of {title}."},
                    {"q": "Electric current is measured in:", "options": ["Amperes", "Volts", "Ohms", "Watts"], "answer": 0, "exp": "The SI unit of electric current is the Ampere (A)."}
                ]
            },
            "englishSummary": f"Unit {num} discusses {title}, examining scientific mechanisms and technological developments for Class 10.",
            "urduSummary": f"یہ یونٹ {titleUr} کے سائنسی اصولوں، روزمرہ زندگی کے تجربات اور ماحولیاتی پہلوؤں کو بیان کرتا ہے۔"
        }
        units.append(unit_obj)
        
    out_file = r"D:\SpaceBook\SpaceBook Web\js\general_science_10_data.js"
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Class 10 General Science Dataset (KPTBB)\n")
        f.write("const GENERAL_SCIENCE_10_DATA = " + json.dumps(units, ensure_ascii=False, indent=2) + ";\n\n")
        f.write("if (typeof DATA !== 'undefined' && DATA) { DATA.genSci10Chapters = GENERAL_SCIENCE_10_DATA; }\n")
        f.write("if (typeof window !== 'undefined') { window.GENERAL_SCIENCE_10_DATA = GENERAL_SCIENCE_10_DATA; }\n")
    print("Generated general_science_10_data.js!")

def build_hpe_10():
    folder = r"D:\SpaceBook\Books\10th\10th HPE\Word"
    paras = []
    for f in sorted(os.listdir(folder)):
        if f.endswith(".docx"):
            doc = docx.Document(os.path.join(folder, f))
            paras.extend([p.text.strip() for p in doc.paragraphs if p.text.strip()])
            
    bounds = [
        (1, "Non-Contagious Diseases (Asthma & Blood Pressure)", "غیر متعدی امراض (دمہ اور فشار خون)", 164, 450),
        (2, "Circulatory System", "نظام دوران خون", 450, 900),
        (3, "Malnutrition & Dietetics", "ناقص غذا اور غذائیت", 900, 1350),
        (4, "Bullying Prevention in Schools", "سکول میں غنڈہ گردی کی روک تھام", 1350, 1800),
        (5, "Emergency Situations & First Aid", "ہنگامی صورت حال اور ابتدائی طبی امداد", 1800, 2250),
        (6, "Warm-up & Cool-down in Athletics", "کھیلوں کے لیے وارم اپ اور وارم ڈاؤن", 2250, 2700),
        (7, "Biomechanics in Sports & Gymnastics", "کھیلوں اور جمناسٹک کے لیے بائیو مکینکس", 2700, 3150),
        (8, "Sports Skills & Tactical Strategies", "کھیلوں کی مہارتیں اور حکمت عملیاں", 3150, len(paras))
    ]
    
    units = []
    for num, title, titleUr, start, end in bounds:
        u_paras = paras[start:end]
        clean_p = [p for p in u_paras if not any(w in p.lower() for w in ["awaz e inqilab", "camscanner", "download pdf", "whatsapp", "not for sale"]) and len(p) > 2]
        sec_paras = clean_p[:40]
        full_text = "\n\n".join(sec_paras[:15])
        
        unit_obj = {
            "number": num,
            "id": f"unit-{num}",
            "title": title,
            "titleUrdu": titleUr,
            "type": "chapter",
            "sections": [
                {
                    "heading": "Textbook Theory & Health Practices",
                    "headingUrdu": "درسی متن و طبی رہنما اصول",
                    "text": full_text,
                    "paras": sec_paras[:12]
                }
            ],
            "exercise": {
                "shortQuestions": [
                    {"q": f"{titleUr} کی تعریف اور بنیادی مقاصد تحریر کریں۔", "a": f"اس باب میں {titleUr} کی تفصیلات، احتیاطی تدابیر اور تندرست زندگی کے اصول بیان کیے گئے ہیں۔"},
                    {"q": f"باب {num} کے دو اہم نکات بیان کریں۔", "a": "1. جسمانی صحت اور باقاعدہ ورزش کی اہمیت۔\n2. احتیاطی تدابیر اور کھیلوں کے دوران حفاظت۔"}
                ],
                "longQuestions": [
                    {"q": f"{titleUr} پر تفصیلی نوٹ تحریر کریں۔", "a": f"اس باب میں {titleUr} کے متعلق سائنسی، طبی اور عملی رہنما اصول جامع انداز میں پیش کیے گئے ہیں۔"}
                ],
                "mcqs": [
                    {"q": f"باب {num} ({titleUr}) کا بنیادی موضوع کیا ہے؟", "options": ["صحت و جسمانی تندرستی", "ریاضی کے کلیے", "کیمیائی مساواتیں", "تاریخی جنگیں"], "answer": 0, "exp": f"یہ باب صحت و جسمانی تعلیم کے تحت {titleUr} کا احاطہ کرتا ہے۔"},
                    {"q": "بلڈ پریشر کی عام نارمل حد کتنی مانی جاتی ہے؟", "options": ["120/80 mmHg", "160/100 mmHg", "90/50 mmHg", "200/120 mmHg"], "answer": 0, "exp": "معمول کے مطابق صحت مند بالغ انسان میں بلڈ پریشر 120/80 mmHg ہوتا ہے۔"}
                ]
            },
            "englishSummary": f"Unit {num} covers {title}, focusing on physiological well-being, injury prevention, and athletic performance.",
            "urduSummary": f"یہ یونٹ {titleUr} کے طبی اصولوں، جسمانی تندرستی اور کھیلوں کی تکنیکوں کی وضاحت کرتا ہے۔"
        }
        units.append(unit_obj)
        
    out_file = r"D:\SpaceBook\SpaceBook Web\js\hpe_10_data.js"
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Class 10 HPE Dataset (KPTBB)\n")
        f.write("const HPE_10_DATA = " + json.dumps(units, ensure_ascii=False, indent=2) + ";\n\n")
        f.write("if (typeof DATA !== 'undefined' && DATA) { DATA.hpe10Chapters = HPE_10_DATA; }\n")
        f.write("if (typeof window !== 'undefined') { window.HPE_10_DATA = HPE_10_DATA; }\n")
    print("Generated hpe_10_data.js!")

def build_drawing_10():
    folder = r"D:\SpaceBook\Books\10th\10th  Drawing\Word"
    paras = []
    for f in sorted(os.listdir(folder)):
        if f.endswith(".docx"):
            doc = docx.Document(os.path.join(folder, f))
            paras.extend([p.text.strip() for p in doc.paragraphs if p.text.strip()])
            
    bounds = [
        (1, "Fundamentals of Drawing, Painting & Design", "ڈرائنگ، پینٹنگ اور ڈیزائن کے بنیادی اصول", 135, 400),
        (2, "Elements of Art, Principles of Design & Artistic Expression", "آرٹ کے عناصر اور ڈیزائن کے اصول", 400, 850),
        (3, "Art Appreciation & Cultural Heritage", "فنون لطیفہ کی تحسین اور ثقافتی ورثہ", 850, 1150),
        (4, "Life Skills Through Art Education", "آرٹ کے ذریعے عملی زندگی کی مہارتیں", 1150, len(paras))
    ]
    
    units = []
    for num, title, titleUr, start, end in bounds:
        u_paras = paras[start:end]
        clean_p = [p for p in u_paras if not any(w in p.lower() for w in ["awaz e inqilab", "camscanner", "download pdf", "whatsapp", "not for sale"]) and len(p) > 2]
        sec_paras = clean_p[:40]
        full_text = "\n\n".join(sec_paras[:15])
        
        unit_obj = {
            "number": num,
            "id": f"unit-{num}",
            "title": title,
            "titleUrdu": titleUr,
            "type": "chapter",
            "sections": [
                {
                    "heading": "Studio Principles & Practical Execution",
                    "headingUrdu": "عملی اصول و فنکارانہ تکنیک",
                    "text": full_text,
                    "paras": sec_paras[:12]
                }
            ],
            "exercise": {
                "shortQuestions": [
                    {"q": f"State the primary concepts of {title}.", "a": f"This chapter explains artistic execution, design guidelines, color balance, and creative tools for {title}."},
                    {"q": f"List essential studio tools mentioned in Chapter {num}.", "a": "Pencils, sketching paper, watercolors, brushes, palette, and geometric drawing instruments."}
                ],
                "longQuestions": [
                    {"q": f"Explain the comprehensive theory and studio practice of {title}.", "a": f"Chapter {num} instructs students on advanced drawing techniques, proportion, composition, and aesthetic judgment."}
                ],
                "mcqs": [
                    {"q": f"What is the primary study of Chapter {num} ({title})?", "options": ["Artistic mediums, design principles, and composition", "Algebraic equations", "Biochemical reactions", "Atomic structure"], "answer": 0, "exp": f"Chapter {num} teaches artistic skills in {title}."},
                    {"q": "Perspective in drawing gives the illusion of:", "options": ["Depth and three-dimensionality", "Only flat lines", "Color inversion", "Sound vibrations"], "answer": 0, "exp": "Perspective provides depth and spatial relationship to 2D artwork."}
                ]
            },
            "englishSummary": f"Chapter {num} covers {title}, offering step-by-step guidance on visual techniques and studio skills.",
            "urduSummary": f"یہ باب {titleUr} کے فنی اصولوں، پینٹنگ تکنیک اور عملی ڈرائنگ کی مہارتوں کی رہنمائی فراہم کرتا ہے۔"
        }
        units.append(unit_obj)
        
    out_file = r"D:\SpaceBook\SpaceBook Web\js\drawing_10_data.js"
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Class 10 Art & Model Drawing Dataset (KPTBB)\n")
        f.write("const DRAWING_10_DATA = " + json.dumps(units, ensure_ascii=False, indent=2) + ";\n\n")
        f.write("if (typeof DATA !== 'undefined' && DATA) { DATA.drawing10Chapters = DRAWING_10_DATA; }\n")
        f.write("if (typeof window !== 'undefined') { window.DRAWING_10_DATA = DRAWING_10_DATA; }\n")
    print("Generated drawing_10_data.js!")

def build_islamiat_ikhtiari_10():
    folder = r"D:\SpaceBook\Books\10th\10th Ismaiat Ikhtiari\Word"
    paras = []
    for f in sorted(os.listdir(folder)):
        if f.endswith(".docx"):
            doc = docx.Document(os.path.join(folder, f))
            paras.extend([p.text.strip() for p in doc.paragraphs if p.text.strip()])
            
    bounds = [
        (1, "Seerat-un-Nabi (PBUH) - Status, Teachings & Uswah", "باب سوم: سیرت رسول اکرم ﷺ (مقام، محبت، اسوہ حسنہ)", 124, 750),
        (2, "Quranic Verses with Translation & Commentary", "باب چہارم: قرآنی آیات (ترجمہ و تشریح)", 750, 1400),
        (3, "Hadith & Sunnah - Compilation, Texts & Ethics", "باب پنجم: حدیث اور سنت (تدوین، کتب، منتخب احادیث)", 1400, 2200),
        (4, "Islamic Sciences & Contributions of Muslim Scholars", "باب ششم: اسلامی علوم اور مسلمانوں کی خدمات", 2200, len(paras))
    ]
    
    units = []
    for num, title, titleUr, start, end in bounds:
        u_paras = paras[start:end]
        clean_p = [p for p in u_paras if not any(w in p.lower() for w in ["awaz e inqilab", "camscanner", "download pdf", "whatsapp", "not for sale"]) and len(p) > 2]
        sec_paras = clean_p[:40]
        full_text = "\n\n".join(sec_paras[:15])
        
        unit_obj = {
            "number": num,
            "id": f"unit-{num}",
            "title": title,
            "titleUrdu": titleUr,
            "type": "chapter",
            "sections": [
                {
                    "heading": "دینی متن و احکامات",
                    "headingUrdu": "دینی متن و احکامات",
                    "text": full_text,
                    "paras": sec_paras[:12]
                }
            ],
            "exercise": {
                "shortQuestions": [
                    {"q": f"{titleUr} کے بنیادی نکات بیان کریں۔", "a": f"اس باب میں {titleUr} کی روشنی میں اسلامی احکامات، اخلاقیات اور عملی زندگی کے تقاضے بیان ہوئے ہیں۔"},
                    {"q": f"باب {num} سے کیا عملی سبق حاصل ہوتا ہے؟", "a": "1. احکام الٰہی کی پابندی اور سنت نبوی ﷺ کی پیروی۔\n2. معاشرے میں عدل، اخوت اور سچائی کو عام کرنا۔"}
                ],
                "longQuestions": [
                    {"q": f"{titleUr} پر جامع و مفصل مضمون لکھیں۔", "a": f"اس باب کا مفصل مطالعہ اسلامی تعلیمات کے مطابق زندگی گزارنے اور سیرت النبی ﷺ کو مشعل راہ بنانے کا شعور عطا کرتا ہے۔"}
                ],
                "mcqs": [
                    {"q": f"باب {num} ({titleUr}) کا بنیادی عنوان کیا ہے؟", "options": ["اسلامی علوم و تعلیمات", "سائنسی ایجادات", "جغرافیائی خدوخال", "لسانی قواعد"], "answer": 0, "exp": f"یہ باب اسلامیات اختیاری کے تحت {titleUr} کی وضاحت کرتا ہے۔"},
                    {"q": "رسول اکرم ﷺ کی اطاعت کا حکم کس نے دیا ہے؟", "options": ["اللہ تعالیٰ نے قرآن مجید میں", "کسی انسان نے نہیں", "صرف ماضی کے لوگوں نے", "اختیاری معاملہ ہے"], "answer": 0, "exp": "قرآن مجید میں اللہ تعالیٰ کا واضح حکم ہے: 'مَنْ يُطِعِ الرَّسُولَ فَقَدْ أَطَاعَ اللَّهَ'۔"}
                ]
            },
            "englishSummary": f"Unit {num} covers {title}, examining the life of the Prophet (PBUH), Hadith literature, and Muslim scholarship.",
            "urduSummary": f"یہ باب {titleUr} کے احکام، مضامین اور اسلامی تعلیمات کی مفصل تشریح پیش کرتا ہے۔"
        }
        units.append(unit_obj)
        
    out_file = r"D:\SpaceBook\SpaceBook Web\js\islamiat_ikhtiari_10_data.js"
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Class 10 Islamiat Ikhtiari Dataset (KPTBB)\n")
        f.write("const ISLAMIAT_IKHTIARI_10_DATA = " + json.dumps(units, ensure_ascii=False, indent=2) + ";\n\n")
        f.write("if (typeof DATA !== 'undefined' && DATA) { DATA.islopt10Chapters = ISLAMIAT_IKHTIARI_10_DATA; }\n")
        f.write("if (typeof window !== 'undefined') { window.ISLAMIAT_IKHTIARI_10_DATA = ISLAMIAT_IKHTIARI_10_DATA; }\n")
    print("Generated islamiat_ikhtiari_10_data.js!")

def build_mutalia_quran_10():
    # Surahs for 10th Mutalia-e-Quran:
    # 1. Surah Al-An'am (سورة الأنعام)
    # 2. Surah Yunus (سورة يونس)
    # 3. Surah Ar-Ra'd (سورة الرعد)
    # 4. Surah Ibrahim (سورة إبراهيم)
    # 5. Surah Al-Hijr (سورة الحجر)
    # 6. Surah An-Nahl (سورة النحل)
    # 7. Surah Bani Isra'il / Al-Isra (سورة بني إسرائيل / الإسراء)
    # 8. Surah Az-Zumar & Al-Mu'min (سورة الزمر وسورة المؤمن / غافر)
    surahs = [
        (1, "Surah Al-An'am (سورة الأنعام)", "سورة الأنعام", "Surah Al-An'am establishes the fundamentals of Tawheed, rejects polytheism, and highlights divine signs in the cosmos.", "سورة الانعام میں توحید باری تعالیٰ، شرک کی تردید اور کائنات میں اللہ کی نشانیوں کا تذکرہ ہے۔"),
        (2, "Surah Yunus (سورة يونس)", "سورة يونس", "Surah Yunus discusses divine revelations, accountability in the Hereafter, and the story of Prophet Yunus (AS).", "سورة یونس میں وحی الٰہی، آخرت میں جوابدہی اور حضرت یونس علیہ السلام کا واقعہ بیان ہوا ہے۔"),
        (3, "Surah Ar-Ra'd (سورة الرعد)", "سورة الرعد", "Surah Ar-Ra'd emphasizes the power of truth, stability of moral law, and the inner peace found in Allah's remembrance.", "سورة الرعد میں حق کی ابدیت، دلوں کے سکون کے لیے ذکر الٰہی اور کائنات کے نظام پر غور کا سبق ہے۔"),
        (4, "Surah Ibrahim (سورة إبراهيم)", "سورة إبراهيم", "Surah Ibrahim focuses on the gratefulness to Allah, the mission of Prophet Ibrahim (AS), and prayer for righteousness.", "سورة ابراہیم میں شکر گزاری کی فضیلت، حضرت ابراہیم علیہ السلام کی دعائیں اور حق پر ثابت قدمی ہے۔"),
        (5, "Surah Al-Hijr (سورة الحجر)", "سورة الحجر", "Surah Al-Hijr assures divine protection for the Quran and recounts lessons from past nations.", "سورة الحجر میں قرآن مجید کی حفاظت کی الٰہی ضمانت اور سابقہ اقوام کے انجام سے عبرت کا درس ہے۔"),
        (6, "Surah An-Nahl (سورة النحل)", "سورة النحل", "Surah An-Nahl describes the abundant blessings of Allah (The Chapter of Blessings) and invites to wisdom and good counsel.", "سورة النحل (سورۃ النعم) میں اللہ کی بے شمار نعمتوں کا شکر اور حکمت کے ساتھ دعوت کا حکم ہے۔"),
        (7, "Surah Bani Isra'il / Al-Isra (سورة بني إسرائيل / الإسراء)", "سورة الإسراء", "Surah Al-Isra details the Miraculous Night Journey (Mi'raj) and delivers universal moral commandments for humanity.", "سورة الاسراء میں واقعہ معراج اور انسانی معاشرے کے لیے جامع اخلاقی و سماجی احکام بیان ہوئے ہیں۔"),
        (8, "Surah Az-Zumar & Al-Mu'min (سورة الزمر وسورة المؤمن)", "سورة الزمر والمؤمن", "Surah Az-Zumar and Al-Mu'min focus on sincere monotheism, repentance, and standing firm in faith.", "سورة الزمر اور سورة المؤمن میں اخلاص فی الدین، توبہ کی اہمیت اور مؤمن آل فرعون کا تاریخی واقعہ ہے۔")
    ]
    
    units = []
    for num, title, titleUr, enSum, urSum in surahs:
        unit_obj = {
            "number": num,
            "id": f"unit-{num}",
            "title": title,
            "titleUrdu": titleUr,
            "type": "chapter",
            "sections": [
                {
                    "heading": "تعارف سورہ مبارکہ و بنیادی مضامین",
                    "headingUrdu": "تعارف سورہ مبارکہ و بنیادی مضامین",
                    "text": f"یہ سورت مبارکہ قرآنی فہم اور ہدایت کا بنیادی ذریعہ ہے۔ اس میں {titleUr} کے بنیادی موضوعات، شان نزول، اخلاقی اسباق اور عقائد کی تفصیلی تشریح موجود ہے۔ طلبہ کو چاہیے کہ وہ آیات کا باقاعدہ مطالعہ کریں، ترجمہ و مفہوم سمجھیں اور ان قرآنی احکام کو اپنی عملی زندگی میں نافذ کریں۔",
                    "paras": [
                        f"{titleUr} کا نام، تعارف اور شان نزول۔",
                        "سورت کے اساسی مضامین: عقیدہ توحید، رسالت اور آخرت۔",
                        "انبیاء کرام کے واقعات اور قوموں کے عروج و زوال کے اسباب۔",
                        "عملی زندگی اور معاشرتی کردار کے لیے قرآنی ہدایات۔"
                    ]
                }
            ],
            "exercise": {
                "shortQuestions": [
                    {"q": f"{titleUr} کے بنیادی مضامین اور مرکزی خیال تحریر کریں۔", "a": f"{titleUr} میں اللہ تعالیٰ کی عظمت، توحید کی حقانیت اور اخلاقی زندگی کے بنیادی اصول بیان کیے گئے ہیں۔"},
                    {"q": f"{titleUr} سے حاصل ہونے والے دو اہم اسباق بیان کریں۔", "a": "1. خالص نیت کے ساتھ اللہ کی عبادت اور اطاعت کرنا۔\n2. ہر حال میں صبر، شکر اور حق کا ساتھ دینا۔"}
                ],
                "longQuestions": [
                    {"q": f"{titleUr} کے پس منظر، بنیادی مضامین اور اہم ہدایات پر تفصیلی روشنی ڈالیں۔", "a": f"اس سورت کا فہم طالب علموں کو قرآنی حکمت، ایمانی پختگی اور اسلامی طرز عمل سے آراستہ کرتا ہے۔"}
                ],
                "mcqs": [
                    {"q": f"{titleUr} کا مرکزی پیغام کیا ہے؟", "options": ["ایمان باللہ، اخلاق اور ہدایت قرآنی", "فقط نحوی قواعد", "تجارتی حساب کتاب", "تاریخی قصے"], "answer": 0, "exp": "قرآن مجید کی تمام سورتیں انسان کی فکری و اخلاقی رہنمائی کا سرچشمہ ہیں۔"},
                    {"q": "قرآنِ حکیم کی تلاوت کے ساتھ ساتھ کیا لازم ہے؟", "options": ["فہم، تدبر اور عملی زندگی میں اتباع", "محض الفاظ سننا", "کوئی عمل نہ کرنا", "بلا فہم آگے بڑھنا"], "answer": 0, "exp": "قرآن کا بنیادی مقصد فہم، تدبر اور عملی زندگی میں نفاذ ہے۔"}
                ]
            },
            "englishSummary": enSum,
            "urduSummary": urSum
        }
        units.append(unit_obj)
        
    out_file = r"D:\SpaceBook\SpaceBook Web\js\mutalia_quran_10_data.js"
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Class 10 Mutalia-e-Quran Dataset (KPTBB)\n")
        f.write("const MUTALIA_QURAN_10_DATA = " + json.dumps(units, ensure_ascii=False, indent=2) + ";\n\n")
        f.write("if (typeof DATA !== 'undefined' && DATA) { DATA.quran10Chapters = MUTALIA_QURAN_10_DATA; }\n")
        f.write("if (typeof window !== 'undefined') { window.MUTALIA_QURAN_10_DATA = MUTALIA_QURAN_10_DATA; }\n")
    print("Generated mutalia_quran_10_data.js!")

if __name__ == "__main__":
    build_cs_10()
    build_gen_science_10()
    build_hpe_10()
    build_drawing_10()
    build_islamiat_ikhtiari_10()
    build_mutalia_quran_10()
    print("All 6 Class 10 datasets generated successfully!")
